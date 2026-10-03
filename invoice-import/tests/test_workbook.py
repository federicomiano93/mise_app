import unittest
from datetime import datetime, timezone
from unittest import mock

from openpyxl import load_workbook

import invoice_import
import workbook
from tests.helpers import TempDirCase

NOW = datetime(2026, 10, 3, 10, 15, tzinfo=timezone.utc)
FLOUR = {"desc": "FARINA TIPO 00 SACCO KG 25", "code": "F00-25", "qty": 25, "unit": "KG", "total": 14.25}
SUGAR = {"desc": "ZUCCHERO SACCHI DA KG 25", "qty": 50, "unit": "KG", "total": 40}
BAGS = {"desc": "BUSTE PLASTICA", "qty": 2, "unit": "PZ", "total": 10}
COLA = {"desc": "COCA COLA LATTINA 33 CL", "qty": 24, "unit": "PZ", "total": 12}


class WorkbookCase(TempDirCase):
    def run_main(self, command, now=NOW):
        messages = []
        code = invoice_import.main([command, "--base", str(self.base)], now=now, out=messages.append)
        return code, "\n".join(messages)

    @property
    def xlsx(self):
        return self.base / "fatture_ingredienti_mise.xlsx"

    def rows(self):
        return {row[workbook.H_KEY]: row for row in workbook.read_products_sheet(self.xlsx)}

    def edit(self, key_part, **cells):
        """Do what the owner does in Excel: change cells of the row whose key contains key_part."""
        wb = load_workbook(self.xlsx)
        ws = wb[workbook.SHEET_PRODUCTS]
        headers = {c.value: c.column for c in ws[workbook.HEADER_ROW]}
        hits = [r for r in range(workbook.HEADER_ROW + 1, ws.max_row + 1)
                if key_part in str(ws.cell(r, headers[workbook.H_KEY]).value)]
        self.assertEqual(len(hits), 1, key_part)
        for header, value in cells.items():
            ws.cell(hits[0], headers[header]).value = value
        wb.save(self.xlsx)
        wb.close()


class LayoutTest(WorkbookCase):
    def test_sheets_headers_and_layout(self):
        self.add([FLOUR, SUGAR, BAGS, COLA])
        code, output = self.run_main("workbook")
        self.assertEqual(code, 0, output)
        wb = load_workbook(self.xlsx)
        self.assertEqual(wb.sheetnames, ["Prodotti", "Righe fattura", "Esclusi", "Fatture"])
        ws = wb["Prodotti"]
        self.assertEqual(workbook.HEADER_ROW, 4)  # three legend rows, then the header
        for r in (1, 2, 3):
            self.assertTrue(ws.cell(r, 1).value)
        header = [c.value for c in ws[workbook.HEADER_ROW]]
        self.assertEqual(header, workbook.HEADERS)
        self.assertEqual(header[0], "Azione")
        self.assertEqual(header[-1], "Chiave (non modificare)")
        self.assertEqual(ws.freeze_panes, "A5")
        self.assertTrue(ws.auto_filter.ref.startswith("A4:"))
        yellow = {h for h, c in zip(header, ws[workbook.HEADER_ROW]) if c.fill.fgColor.rgb.endswith("FFF2CC")}
        self.assertEqual(yellow, set(workbook.EDITABLE))
        wb.close()

    def test_one_row_per_product_sorted_by_type_supplier_name(self):
        self.add([COLA, BAGS, SUGAR, FLOUR])
        self.run_main("workbook")
        wb = load_workbook(self.xlsx)
        ws = wb["Prodotti"]
        col = workbook.HEADERS.index(workbook.H_TYPE) + 1
        types = [ws.cell(r, col).value for r in range(5, ws.max_row + 1)]
        self.assertEqual(types, ["ingrediente", "ingrediente", "packaging", "rivendita"])
        names = [ws.cell(r, workbook.HEADERS.index(workbook.H_NAME) + 1).value for r in range(5, 9)]
        self.assertEqual(names, ["Farina tipo 00", "Zucchero", "Buste plastica", "Coca cola lattina"])
        wb.close()

    def test_proposals(self):
        self.add([FLOUR, SUGAR])
        self.run_main("workbook")
        rows = self.rows()
        flour = rows["IT00000000001|code:F00-25"]
        self.assertEqual(flour[workbook.H_PRICE_UNIT], "kg")
        self.assertEqual(flour[workbook.H_WEIGHT], "25 kg")
        self.assertEqual(flour[workbook.H_PACK], "sacco")
        self.assertEqual(flour[workbook.H_LAST], 0.57)
        self.assertEqual(flour[workbook.H_RELIABILITY], "alta")
        self.assertEqual(flour[workbook.H_ACTION] or "", "")
        self.assertEqual(flour[workbook.H_PURCHASES], 1)

    def test_price_history_columns(self):
        self.add([{**FLOUR, "total": 12.5}], date="2026-09-01")
        self.add([{**FLOUR, "total": 15.0}], date="2026-09-15")
        self.run_main("workbook")
        row = self.rows()["IT00000000001|code:F00-25"]
        self.assertEqual((row[workbook.H_LAST], row[workbook.H_MIN], row[workbook.H_MAX]), (0.6, 0.5, 0.6))
        self.assertEqual(row[workbook.H_VARIATION], 20.0)
        self.assertEqual(row[workbook.H_PURCHASES], 2)
        self.assertEqual(row[workbook.H_LAST_DATE], "2026-09-15")

    def test_other_sheets_list_lines_exclusions_and_documents(self):
        self.add([FLOUR, {"desc": "LOTTO 99"}])
        self.add([FLOUR], doc_type="TD04", number="2", date="2026-09-02")
        self.run_main("workbook")
        wb = load_workbook(self.xlsx)
        self.assertEqual(wb["Righe fattura"].max_row, 2)  # header + the one kept line
        excluded = [[c.value for c in row] for row in wb["Esclusi"].iter_rows(min_row=2)]
        self.assertEqual(sorted(r[7] for r in excluded), ["riga informativa", "tipo documento TD04"])
        docs = list(wb["Fatture"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(docs), 2)
        self.assertEqual({d[3] for d in docs}, {"TD01", "TD04"})
        wb.close()

    def test_formula_text_in_a_description_stays_text(self):
        self.add([{"desc": "=HYPERLINK(\"http://example.invalid\")", "qty": 1, "unit": "KG", "total": 3}])
        self.run_main("workbook")
        wb = load_workbook(self.xlsx)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", f"{ws.title}!{cell.coordinate}")
        wb.close()

    def test_a_supplier_without_vat_never_shows_its_tax_code_anywhere(self):
        tax_code = "RSSMRA80A01H501U"
        self.add([FLOUR], name="Mario Rossi", vat="", tax_code=tax_code, base=f"{tax_code}_00001")
        self.run_main("workbook")
        wb = load_workbook(self.xlsx)
        seen = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
        wb.close()
        self.assertTrue(any(v.startswith("NOVAT:") for v in seen))
        self.assertFalse([v for v in seen if tax_code.lower() in v.lower()])

    def test_summary_is_printed(self):
        self.add([FLOUR, BAGS])
        code, output = self.run_main("workbook")
        self.assertIn("Prodotti: 2", output)
        self.assertIn("ingrediente: 1", output)
        self.assertIn("da verificare:", output)


class CarryOverTest(WorkbookCase):
    def test_decisions_survive_a_second_run_and_the_old_file_is_backed_up(self):
        self.add([FLOUR, BAGS])
        self.run_main("workbook")
        self.edit("code:F00-25", **{
            workbook.H_ACTION: "importa", workbook.H_NAME: "Farina del forno", workbook.H_BRAND: "Marca Esempio",
            workbook.H_CATEGORY: "Farine", workbook.H_WEIGHT: "20 kg", workbook.H_PACK: "sacchetto",
        })
        self.edit("name:buste", **{workbook.H_ACTION: "scarta"})
        self.add([{**FLOUR, "total": 15}], date="2026-09-20")  # a new month's invoice
        code, output = self.run_main("workbook", now=datetime(2026, 11, 1, 9, 30, tzinfo=timezone.utc))
        self.assertEqual(code, 0, output)
        self.assertIn("riportate dal foglio precedente: 2", output)  # only real decisions are counted
        rows = self.rows()
        flour = rows["IT00000000001|code:F00-25"]
        self.assertEqual(flour[workbook.H_ACTION], "importa")
        self.assertEqual(flour[workbook.H_NAME], "Farina del forno")
        self.assertEqual(flour[workbook.H_BRAND], "Marca Esempio")
        self.assertEqual(flour[workbook.H_CATEGORY], "Farine")
        self.assertEqual(flour[workbook.H_WEIGHT], "20 kg")
        self.assertEqual(flour[workbook.H_PACK], "sacchetto")
        self.assertEqual(flour[workbook.H_PURCHASES], 2)  # the data is fresh
        self.assertEqual(flour[workbook.H_LAST], 0.6)
        backups = sorted(p.name for p in self.base.glob("fatture_ingredienti_mise.*.xlsx"))
        self.assertEqual(backups, ["fatture_ingredienti_mise.2026-11-01_0930.xlsx"])
        old = {r[workbook.H_KEY]: r for r in workbook.read_products_sheet(self.base / backups[0])}
        self.assertEqual(old["IT00000000001|code:F00-25"][workbook.H_ACTION], "importa")

    def test_empty_cells_do_not_wipe_proposals(self):
        self.add([FLOUR])
        self.run_main("workbook")
        self.edit("code:F00-25", **{workbook.H_NAME: None, workbook.H_PACK: None})
        self.run_main("workbook")
        row = self.rows()["IT00000000001|code:F00-25"]
        self.assertEqual(row[workbook.H_NAME], "Farina tipo 00")
        self.assertEqual(row[workbook.H_PACK], "sacco")

    def test_a_corrected_weight_changes_the_prices_shown(self):
        self.add([{"desc": "SPEZIA MISTA", "qty": 4, "unit": "PZ", "total": 20}])
        self.run_main("workbook")
        key = "name:spezia mista"
        self.assertIn(self.rows()[next(k for k in self.rows() if key in k)][workbook.H_LAST], (None, ""))
        self.edit(key, **{workbook.H_PRICE_UNIT: "kg", workbook.H_WEIGHT: "2,5 kg"})
        self.run_main("workbook")
        self.assertEqual(self.rows()[next(k for k in self.rows() if key in k)][workbook.H_LAST], 2.0)

    def test_two_runs_in_one_minute_keep_both_backups(self):
        self.add([FLOUR])
        for _ in range(3):
            self.run_main("workbook")
        self.assertEqual(len(list(self.base.glob("fatture_ingredienti_mise.*.xlsx"))), 2)


class LockedFileTest(WorkbookCase):
    def test_a_workbook_open_in_excel_is_left_untouched(self):
        self.add([FLOUR])
        self.run_main("workbook")
        before = self.xlsx.read_bytes()
        with mock.patch.object(workbook, "probe_writable", side_effect=PermissionError("locked")):
            code, output = self.run_main("workbook")
        self.assertEqual(code, 1)
        self.assertIn("Chiudi il file in Excel e riprova", output)
        self.assertEqual(self.xlsx.read_bytes(), before)
        self.assertEqual(list(self.base.glob("fatture_ingredienti_mise.*.xlsx")), [])
        self.assertEqual([p.name for p in self.base.iterdir() if p.suffix == ".tmp"], [])

    def test_a_lock_discovered_while_swapping_in_the_new_file_is_friendly_too(self):
        self.add([FLOUR])
        self.run_main("workbook")
        before = self.xlsx.read_bytes()
        with mock.patch.object(workbook.os, "replace", side_effect=PermissionError("locked")):
            code, output = self.run_main("workbook")
        self.assertEqual(code, 1)
        self.assertIn("Chiudi il file in Excel", output)
        self.assertEqual(self.xlsx.read_bytes(), before)
        self.assertEqual([p.name for p in self.base.iterdir() if p.name.endswith(".tmp")], [])


if __name__ == "__main__":
    unittest.main()
