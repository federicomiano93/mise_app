"""The decisions workbook: one row per product, the owner fills «Azione» and corrects
the yellow columns. A second run keeps what was typed (matched by the «Chiave» column).
"""
from __future__ import annotations

import os
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from classify import format_pack_size
from pricing import Params, compute_document, round_half_up
from products import (
    TYPE_ORDER, Catalogue, Product, evaluate, params_from_cells, propose_name, propose_pack_word, propose_params,
)

SHEET_PRODUCTS = "Prodotti"

H_ACTION = "Azione"
H_TYPE = "Tipo"
H_SUPPLIER = "Fornitore"
H_VAT = "P.IVA"
H_DESCRIPTION = "Descrizione in fattura (ultima)"
H_CODE = "Codice articolo"
H_NAME = "Nome in Mise"
H_BRAND = "Marca"
H_CATEGORY = "Categoria"
H_PRICE_UNIT = "Unità prezzo (kg/l/pz)"
H_WEIGHT = "Peso confezione"
H_PACK = "Confezione"
H_COUNT = "Pezzi per cartone"
H_LAST = "Ultimo prezzo €/unità netto"
H_MIN = "Min €/unità"
H_MAX = "Max €/unità"
H_VARIATION = "Variazione %"
H_VAT_RATE = "IVA %"
H_PURCHASES = "N. acquisti"
H_QTY = "Quantità totale"
H_LAST_DATE = "Ultimo acquisto"
H_RELIABILITY = "Affidabilità"
H_NOTE = "Nota"
H_KEY = "Chiave (non modificare)"

HEADERS = [
    H_ACTION, H_TYPE, H_SUPPLIER, H_VAT, H_DESCRIPTION, H_CODE, H_NAME, H_BRAND, H_CATEGORY, H_PRICE_UNIT,
    H_WEIGHT, H_PACK, H_COUNT, H_LAST, H_MIN, H_MAX, H_VARIATION, H_VAT_RATE, H_PURCHASES, H_QTY, H_LAST_DATE,
    H_RELIABILITY, H_NOTE, H_KEY,
]
# The columns the owner edits. A second run copies their non-empty values over.
EDITABLE = [H_ACTION, H_TYPE, H_NAME, H_BRAND, H_CATEGORY, H_PRICE_UNIT, H_WEIGHT, H_PACK, H_COUNT]

LEGEND = [
    "Compila la colonna «Azione»: «importa» porta il prodotto in Mise; «unisci con <nome>» lo unisce a un "
    "ingrediente già presente in Mise (scrivi il nome esatto); «scarta» o vuoto lo lascia fuori.",
    "Le colonne gialle sono tue: correggile pure, restano anche quando rigeneri il foglio. Non toccare «Chiave». "
    "I prezzi sono netti IVA, per unità prezzo (kg, l o pezzo). Tipo: ingrediente, packaging o rivendita — per ora "
    "si importano solo gli ingredienti (packaging e rivendita restano fuori dal file).",
    "Peso confezione = UNA confezione («2.5 kg», «250 g», «1 l», «500 ml»); Pezzi per cartone = quante confezioni "
    "contiene il cartone. Affidabilità: alta = fatturato a peso o volume; media = peso letto dalla descrizione o "
    "uova; da verificare = controlla a mano. Variazione % = dal primo all'ultimo acquisto.",
]
HEADER_ROW = len(LEGEND) + 1

YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="EDEDED")
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")


class WorkbookError(Exception):
    """Something the owner can fix; the message is shown as it is."""


class WorkbookLocked(WorkbookError):
    pass


# ── reading ──────────────────────────────────────────────────────────────────


def read_products_sheet(path: Path) -> list[dict]:
    """Every product row of an existing workbook as {header: value}, found BY HEADER NAME
    (the owner may add or move columns, and the legend above the header can change)."""
    try:
        wb = load_workbook(path, data_only=True)
    except PermissionError as exc:
        raise WorkbookLocked(str(exc)) from exc
    except Exception as exc:  # a damaged file raises many different things
        raise WorkbookError(f"Non riesco a leggere {path.name}: {exc}") from exc
    try:
        if SHEET_PRODUCTS not in wb.sheetnames:
            raise WorkbookError(f"Nel file {path.name} manca il foglio «{SHEET_PRODUCTS}».")
        ws = wb[SHEET_PRODUCTS]
        header_row = None
        headers: list = []
        for r in range(1, min(ws.max_row, 30) + 1):
            values = [c.value for c in ws[r]]
            if H_ACTION in values and H_KEY in values:
                header_row, headers = r, values
                break
        if header_row is None:
            raise WorkbookError(f"Nel foglio «{SHEET_PRODUCTS}» non trovo la riga delle intestazioni.")
        rows = []
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            row = {h: v for h, v in zip(headers, values) if h}
            if str(row.get(H_KEY) or "").strip():
                row[H_KEY] = str(row[H_KEY]).strip()
                rows.append(row)
        return rows
    finally:
        wb.close()


# ── building the rows ────────────────────────────────────────────────────────


@dataclass
class ProductRow:
    product: Product
    cells: dict
    params: Params
    sort_key: tuple
    decided: bool = False


def unit_cell(price_unit: str | None) -> str:
    return {"pcs": "pz"}.get(price_unit or "", price_unit or "")


def build_rows(catalogue: Catalogue, carried: dict[str, dict]) -> list[ProductRow]:
    rows = []
    for product in catalogue.products.values():
        proposal = propose_params(product, catalogue.config)
        cells = {
            H_ACTION: "",
            H_TYPE: product.type,
            H_NAME: propose_name(product),
            H_BRAND: "",
            H_CATEGORY: "",
            H_PRICE_UNIT: unit_cell(proposal.price_unit),
            H_WEIGHT: format_pack_size(proposal.pack.size, proposal.pack.unit) if proposal.pack else "",
            H_PACK: propose_pack_word(product),
            H_COUNT: proposal.pack_count or "",
        }
        decided = False  # a carried value that differs from the proposal is a decision of the owner
        for header, value in carried.get(product.key, {}).items():
            if header in EDITABLE and value is not None and str(value).strip() != "":
                decided = decided or str(value).strip() != str(cells[header]).strip()
                cells[header] = value
        params, problems = params_from_cells(cells[H_PRICE_UNIT], cells[H_WEIGHT], cells[H_COUNT], proposal)
        ev = evaluate(product, params, catalogue, problems)
        supplier = catalogue.suppliers[product.supplier_key]
        points = ev.points
        latest_line = product.latest.lines[0]
        cells.update({
            H_SUPPLIER: supplier.name,
            H_VAT: supplier.vat_number,
            H_DESCRIPTION: product.description,
            H_CODE: product.code,
            H_LAST: points[-1].price if points else "",
            H_MIN: min(p.price for p in points) if points else "",
            H_MAX: max(p.price for p in points) if points else "",
            H_VARIATION: (
                round_half_up((points[-1].price - points[0].price) / points[0].price * 100, 1)
                if len(points) > 1 and points[0].price > 0 else ""
            ),
            H_VAT_RATE: latest_line.vat_rate if latest_line.vat_rate is not None else "",
            H_PURCHASES: len(product.entries),
            H_QTY: round_half_up(sum(p.qty for p in points), 3) if points else "",
            H_LAST_DATE: product.latest.doc.date,
            H_RELIABILITY: ev.reliability,
            H_NOTE: ev.note,
            H_KEY: product.key,
        })
        sort_key = (TYPE_ORDER[product.type], supplier.name.casefold(), str(cells[H_NAME]).casefold(), product.key)
        rows.append(ProductRow(product, cells, params, sort_key, decided))
    rows.sort(key=lambda r: r.sort_key)
    return rows


# ── writing ──────────────────────────────────────────────────────────────────

WIDTHS = {
    H_ACTION: 16, H_TYPE: 13, H_SUPPLIER: 28, H_VAT: 15, H_DESCRIPTION: 46, H_CODE: 14, H_NAME: 30,
    H_BRAND: 14, H_CATEGORY: 14, H_PRICE_UNIT: 12, H_WEIGHT: 14, H_PACK: 13, H_COUNT: 11, H_LAST: 14,
    H_MIN: 11, H_MAX: 11, H_VARIATION: 11, H_VAT_RATE: 8, H_PURCHASES: 10, H_QTY: 13, H_LAST_DATE: 13,
    H_RELIABILITY: 14, H_NOTE: 44, H_KEY: 40,
}
NUMBER_FORMATS = {
    H_LAST: "0.0000", H_MIN: "0.0000", H_MAX: "0.0000", H_VARIATION: "0.0", H_QTY: "#,##0.000",
    H_VAT_RATE: "0", H_PURCHASES: "0",
}


def _put(ws, row: int, col: int, value):
    """Write a cell. Text from an invoice that starts with «=» must stay text: openpyxl
    would otherwise store it as a formula and Excel would run it."""
    cell = ws.cell(row, col, value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def _header_cell(cell, fill) -> None:
    cell.font = Font(bold=True)
    cell.fill = fill
    cell.alignment = Alignment(wrap_text=True, vertical="center")


def _products_sheet(ws, rows: list[ProductRow]) -> None:
    last_col = get_column_letter(len(HEADERS))
    for i, text in enumerate(LEGEND, start=1):
        ws.merge_cells(f"A{i}:{last_col}{i}")
        ws.cell(i, 1, text).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 32
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(HEADER_ROW, col, header)
        _header_cell(cell, YELLOW if header in EDITABLE else GREY if header == H_KEY else HEADER_FILL)
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[header]
    ws.row_dimensions[HEADER_ROW].height = 32
    for r, row in enumerate(rows, start=HEADER_ROW + 1):
        for col, header in enumerate(HEADERS, start=1):
            cell = _put(ws, r, col, row.cells[header])
            if header in EDITABLE:
                cell.fill = YELLOW
            elif header == H_KEY:
                cell.fill = GREY
            if header in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[header]
            if header in (H_NOTE, H_DESCRIPTION):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 1)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{max(HEADER_ROW + len(rows), HEADER_ROW + 1)}"
    if rows:
        # A hint, not a rule: «unisci con …» is free text.
        action = DataValidation(type="list", formula1='"importa,scarta"', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(action)
        action.add(f"A{HEADER_ROW + 1}:A{HEADER_ROW + len(rows)}")
        kind = DataValidation(type="list", formula1=f'"{",".join(TYPE_ORDER)}"', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(kind)
        type_col = get_column_letter(HEADERS.index(H_TYPE) + 1)
        kind.add(f"{type_col}{HEADER_ROW + 1}:{type_col}{HEADER_ROW + len(rows)}")


def _plain_sheet(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    for col, header in enumerate(headers, start=1):
        _header_cell(ws.cell(1, col, header), HEADER_FILL)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    for r, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            _put(ws, r, col, value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"


def _lines_sheet(ws, rows: list[ProductRow], catalogue: Catalogue) -> None:
    out = []
    for prow in rows:
        for entry in prow.product.entries:
            for line in entry.lines:
                one = compute_document([line], prow.params)
                out.append([
                    prow.cells[H_SUPPLIER], prow.cells[H_VAT], entry.doc.number, entry.doc.date, entry.doc.sdi_id,
                    line.number, line.article_code, line.description, line.quantity, line.unit, line.unit_price,
                    line.total, line.vat_rate, line.natura, "sì" if line.has_discount else "",
                    entry.type, one.price if one.price is None else round_half_up(one.price, 4),
                    unit_cell(prow.params.price_unit), None if one.qty is None else round_half_up(one.qty, 3),
                    one.reliability, one.note, prow.product.key, entry.doc.file,
                ])
    out.sort(key=lambda v: (str(v[0]).casefold(), v[3], str(v[2]), v[5]))
    headers = ["Fornitore", "P.IVA", "Fattura n.", "Data", "Id SDI", "Linea", "Codice articolo", "Descrizione",
               "Quantità", "UM", "Prezzo unitario", "Prezzo totale", "IVA %", "Natura", "Sconto sulla riga",
               "Tipo proposto", "Prezzo €/unità (riga)", "Unità prezzo", "Quantità in unità prezzo", "Affidabilità",
               "Nota", "Chiave", "File"]
    _plain_sheet(ws, headers, out, [28, 15, 12, 12, 12, 7, 14, 46, 10, 7, 12, 12, 7, 8, 9, 12, 13, 9, 13, 14, 40, 40, 30])


def _excluded_sheet(ws, catalogue: Catalogue) -> None:
    rows = [[e.level, e.supplier, e.vat_number, e.number, e.date, e.line, e.description, e.reason, e.file]
            for e in catalogue.excluded]
    _plain_sheet(ws, ["Livello", "Fornitore", "P.IVA", "Fattura n.", "Data", "Linea", "Descrizione", "Motivo", "File"],
                 rows, [11, 28, 15, 12, 12, 7, 50, 40, 30])


def _documents_sheet(ws, catalogue: Catalogue) -> None:
    rows = [[d.file, d.supplier_name, d.vat_number, d.doc_type, d.date, d.number, d.total, d.sdi_id,
             d.reception_date, catalogue.doc_status.get(d.doc_id, "")] for d in catalogue.documents]
    _plain_sheet(ws, ["File", "Fornitore", "P.IVA", "Tipo", "Data", "Numero", "Totale", "Id SDI", "Data ricezione",
                      "Esito"], rows, [34, 28, 15, 7, 12, 12, 12, 12, 14, 40])


def build_workbook(catalogue: Catalogue, rows: list[ProductRow]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PRODUCTS
    _products_sheet(ws, rows)
    _lines_sheet(wb.create_sheet("Righe fattura"), rows, catalogue)
    _excluded_sheet(wb.create_sheet("Esclusi"), catalogue)
    _documents_sheet(wb.create_sheet("Fatture"), catalogue)
    return wb


def _backup_path(path: Path, now: datetime) -> Path:
    stamp = now.strftime("%Y-%m-%d_%H%M")
    candidate = path.with_name(f"{path.stem}.{stamp}{path.suffix}")
    n = 2
    while candidate.exists():  # two runs in one minute must not overwrite the first copy
        candidate = path.with_name(f"{path.stem}.{stamp}-{n}{path.suffix}")
        n += 1
    return candidate


def probe_writable(path: Path) -> None:
    """Excel locks an open workbook: opening it for append fails with PermissionError and
    changes nothing if it succeeds."""
    with open(path, "ab"):
        pass


def write_workbook(path: Path, catalogue: Catalogue, now: datetime) -> dict:
    """Carry the owner's decisions over, back the old file up, write the new one.

    Nothing is touched until the file is known to be writable; the new file is written
    beside the old one and swapped in, so a crash never leaves half a workbook.
    """
    carried: dict[str, dict] = {}
    if path.exists():
        try:
            probe_writable(path)
        except PermissionError as exc:
            raise WorkbookLocked(str(exc)) from exc
        carried = {row[H_KEY]: row for row in read_products_sheet(path)}
    rows = build_rows(catalogue, carried)
    wb = build_workbook(catalogue, rows)
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(tmp)
        if path.exists():
            shutil.copy2(path, _backup_path(path, now))
        os.replace(tmp, path)
    except PermissionError as exc:
        raise WorkbookLocked(str(exc)) from exc
    finally:
        if tmp.exists():
            tmp.unlink()
    return summarise(rows, catalogue, carried)


def summarise(rows: list[ProductRow], catalogue: Catalogue, carried: dict) -> dict:
    by_type: dict[str, int] = {}
    by_reliability: dict[str, int] = {}
    for r in rows:
        by_type[r.cells[H_TYPE]] = by_type.get(r.cells[H_TYPE], 0) + 1
        by_reliability[r.cells[H_RELIABILITY]] = by_reliability.get(r.cells[H_RELIABILITY], 0) + 1
    return {
        "documents": sum(1 for s in catalogue.doc_status.values() if s == "inclusa"),
        "documentsExcluded": sum(1 for s in catalogue.doc_status.values() if s != "inclusa"),
        "products": len(rows),
        "byType": by_type,
        "byReliability": by_reliability,
        "carriedOver": sum(1 for r in rows if r.decided),
    }
